from openpyxl import load_workbook
from sqlmodel import Session, select
from app.models import Equipement
from datetime import date

FEUILLES_A_IGNORER = [
    ' ORDINATEURS ACV',
    'LOCALISATION EQUIPEMENTS RESEAU',
    'SYNTHESE',
    'Feuil1',
    'DOMOTIQUE G',
    'SERVICE ADSL FIBRE',
]

def get_col(row, index):
    try:
        return row[index]
    except IndexError:
        return None

def import_excel(filepath: str, session: Session):
    workbook = load_workbook(filename=filepath)
    for sheet_name in workbook.sheetnames:
        if sheet_name in FEUILLES_A_IGNORER:
            continue  
        sheet = workbook[sheet_name]
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            # 1. SÉCURITÉ : On vérifie si la ligne est vide ou n'a pas de contenu exploitable
            if not row or all(v is None for v in row):
                continue
                
            # 2. LE BANNING : Si une cellule contient une formule (=COUNTIF...), on vire toute la ligne
            # On vérifie les 3 premières cellules par sécurité
            is_formula = any(isinstance(cell, str) and cell.strip().startswith("=") for cell in row[:3])
            if is_formula:
                print(f" Ligne de statistique bannie avec succès dans la feuille : {sheet_name}")
                continue # On passe directement à la ligne suivante de l'Excel
            
            # 3. EXTRACTION (Alignée correctement dans la boucle 'for row')
            salle = get_col(row, 0)
            categorie = get_col(row, 1)
            date_installation = get_col(row, 2)
            nom = get_col(row, 3)
            marque = get_col(row, 4)
            modele = get_col(row, 5)
            numero_serie = get_col(row, 6)
            version_os = get_col(row, 7)
            commentaire = get_col(row, 8)
            cout = get_col(row, 9)

            if not numero_serie:
                continue

            # Nettoyage rapide du numéro de série s'il est récupéré bizarrement
            numero_serie_str = str(numero_serie).strip()
            if numero_serie_str.startswith("="):
                continue

            existing = session.exec(select(Equipement).where(Equipement.numero_serie == numero_serie_str)).first()
            if existing:
                continue

            if date_installation:
                try:
                    date_installation = date(int(str(date_installation).strip()), 1, 1)
                except (ValueError, TypeError):
                    date_installation = None
            else:
                date_installation = None

            equipement = Equipement(
                salle=salle or "",
                categorie=categorie or "",
                date_installation=date_installation,
                nom=nom or "",
                marque=marque or "",
                modele=modele or "",
                numero_serie=numero_serie_str,
                version_os=version_os,
                commentaire=commentaire,
                cout=str(cout) if cout else None,
                etat="fonctionnel",
                usager="",
                lieu=sheet_name,
            )
            session.add(equipement)
            
        # On commit à la fin de chaque feuille pour valider les données sereinement
        session.commit()