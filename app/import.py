from app.dependencies import engine
from sqlmodel import Session
from app.utils.excel_parser import import_excel
from openpyxl import load_workbook

wb = load_workbook("app/database/inventaire.xlsx", data_only=True)
print(wb.sheetnames)

with Session(engine) as session:
    import_excel("app/database/inventaire.xlsx", session)
    print("Import terminé")